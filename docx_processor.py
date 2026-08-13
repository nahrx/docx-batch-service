import sys
import json
from docx import Document
from docx.oxml.ns import qn
from copy import deepcopy

def include_blocks(doc, included_blocks):
    body = doc.element.body
    remove_mode = False
    current_end = None

    for element in list(body):
        # Paragraph
        if element.tag == qn('w:p'):
            text = "".join(
                t.text for t in element.iter(qn('w:t')) if t.text
            ).strip()

            if not remove_mode:
                for block in included_blocks:
                    start, end, included = block[0],block[1],block[2]
                    if included :
                        continue
                    if text == start:
                        remove_mode = True
                        current_end = end
                        body.remove(element)
                        break
            else:
                if text == current_end:
                    remove_mode = False
                    current_end = None
                body.remove(element)

        # Table
        elif element.tag == qn('w:tbl'):
            if remove_mode:
                body.remove(element)

# def replace_in_paragraph(p, replacements):
#     if not p.runs:
#         return

#     full_text = "".join(run.text for run in p.runs)

#     new_text = full_text
#     for key, value in replacements.items():
#         new_text = new_text.replace(key, value)

#     if new_text != full_text:
#         for run in p.runs:
#             run.text = ""
#         p.runs[0].text = new_text

def replace_in_paragraph(p, replacements):
    if not p.runs:
        return

    # 1. Collect full text and per-character formatting
    full_text = ""
    char_styles = []  # one entry per character

    for run in p.runs:
        run_text = run.text
        full_text += run_text
        for ch in run_text:
            char_styles.append({
                "char":ch,
                "bold": run.bold,
                "italic": run.italic,
                "underline": run.underline,
                "font": run.font.name,
                "size": run.font.size
            })

    # 2. Replace placeholders in text
    new_text = full_text

    replacements_log = []
    for key, value in replacements.items():
        start = 0
        while True:
            idx = new_text.find(key, start)
            if idx == -1:
                break

            replacements_log.append({
                "key": key,
                "value": value,
                "start": idx,
                "end": idx + len(key)
            })

            # move search index forward
            start = idx + len(key)
        new_text = new_text.replace(key, value)

    if new_text == full_text:
        return

    # 3. Rebuild paragraph preserving formatting
    p.clear()

    new_char_styles = apply_style_replacements(char_styles, replacements_log)
    i = 0
    while i < len(new_text):
        # Find which original style to use
        if i < len(new_char_styles):
            style = new_char_styles[i]
        else:
            style = new_char_styles[-1]

        run = p.add_run(new_text[i])
        run.bold = style["bold"]
        run.italic = style["italic"]
        run.underline = style["underline"]
        run.font.name = style["font"]
        run.font.size = style["size"]

        i += 1

# def apply_style_replacements(styles, replacement_log):
#     new_styles = styles[:]

#     # process from back to front
#     print(replacement_log)
#     for r in sorted(replacement_log, key=lambda x: x["start"], reverse=True):
#         start = r["start"]
#         end = r["end"]
#         value_len = len(r["value"])
#         # print(r["value"])
#         # print(r["start"])
#         # print(r["end"])
#         # print(len(new_styles))

#         base_style = new_styles[start]

#         # replace slice
#         new_styles[start:end] = [base_style] * value_len

#     return new_styles

def apply_style_replacements(styles, replacement_log):
    new_styles = styles[:]

    for r in sorted(replacement_log, key=lambda x: x["start"]):
        start = r["start"]
        end = r["end"]
        value_len = len(r["value"])

        base_style = new_styles[start]

        new_styles[start:end] = [base_style] * value_len


    return new_styles

def replace_placeholders(doc, replacements):
    # 1️⃣ Normal document paragraphs
    for p in doc.paragraphs:
        replace_in_paragraph(p, replacements)

    # 2️⃣ Table paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, replacements)

def build_combined_replacements(replacements, included_blocks):
    """
    replacements: dict[str, str]
    included_blocks: list of [start, end, bool]

    returns: dict[str, str]
    """

    combined = dict(replacements)  # copy base replacements

    for start, end, _ in included_blocks:
        combined[start] = ""
        combined[end] = ""

    return combined

def fix_footer_first_page_only(doc):
    for i, section in enumerate(doc.sections):
        section.different_first_page_header_footer = True

        # VERY IMPORTANT with section breaks
        section.footer.is_linked_to_previous = False
        section.first_page_footer.is_linked_to_previous = False

        # Clear footers for ALL sections except first
        if i > 0:
            for el in list(section.footer._element):
                section.footer._element.remove(el)
            for el in list(section.first_page_footer._element):
                section.first_page_footer._element.remove(el)

def replace_in_cell(cell, replacements):
    if not cell.paragraphs:
        return

    # Merge all paragraphs + runs
    merged_text = "\n".join(
        "".join(run.text for run in p.runs)
        for p in cell.paragraphs
    )

    # Clear cell completely
    cell._tc.clear_content()

    # Create a single paragraph
    p = cell.add_paragraph()
    p.add_run(merged_text)

    # Replace using your trusted function
    replace_in_paragraph(p, replacements)

def insert_tables_from_markers(doc, tables):
    for table in doc.tables:
        for row in table.rows:
            row_text = "".join(cell.text for cell in row.cells)
            

            for table_marker, table_data in tables.items():
                
                if table_marker not in row_text:
                    continue
                
                # how many data rows
                num_rows = max(len(v) for v in table_data.values())
                

                template_tr = row._tr
                tbl = table._tbl
                insert_at = tbl.index(template_tr)
                for i in range(num_rows):
                    new_tr = deepcopy(template_tr)
                    tbl.insert(insert_at + i, new_tr)

                    new_row = table.rows[insert_at+i-2]
                    print(insert_at,i)
                
                    for cell in new_row.cells:
                        for p in cell.paragraphs:
                            if not p.runs:
                                continue
                            
                            full_text = "".join(run.text for run in p.runs)
                            new_text = full_text
                            
                            for marker, value in table_data.items():
                                new_text = new_text.replace(marker, value[i])
                                new_text = new_text.replace(table_marker, "")

                            if new_text != full_text:
                                for run in p.runs:
                                    run.text = ""
                                p.runs[0].text = new_text

                # remove template row
                tbl.remove(template_tr)
                return
def replace_placeholders_in_headers_and_footers(doc, replacements):
    for section in doc.sections:
        # Headers
        for header in [
            section.header,
            section.first_page_header
        ]:
            for p in header.paragraphs:
                replace_in_paragraph(p, replacements)
            for table in header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            replace_in_paragraph(p, replacements)

        # Footers
        for footer in [
            section.footer,
            section.first_page_footer
        ]:
            for p in footer.paragraphs:
                replace_in_paragraph(p, replacements)
            for table in footer.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            replace_in_paragraph(p, replacements)
def test_insert_tables():
    # 1️⃣ Load template
    doc = Document("template/SPK_template.docx")  # <-- your Word file

    # 2️⃣ Test JSON
    tables = {
        "{{T1}}": {
            "{{No}}": ["1", "2", "3"],
            "{{JabatanKegiatan}}": ["PPL", "PPL", "PML"],
            "{{Kegiatan}}": [
                "Pendataan Survei Ekonomi Nasional 2026",
                "Pendataan Survei Ekonomi Rumah Tangga TW I 2026",
                "Pendataan Survei HK Februari 2026"
            ],
            "{{Volume}}": ["20", "10", "30"],
            "{{Satuan}}": ["RT", "RT", "Dokumen"],
            "{{JadwalKegiatan}}": [
                "1–20 Feb 2026",
                "5–15 Feb 2026",
                "25 Feb – 5 Mar 2026"
            ],
            "{{HonorSatuan}}": ["27000", "17000", "9000"],
            "{{HonorKegiatan}}": ["540000", "170000", "270000"],
            "{{MAK}}": [
                "2906.BMA.006.005.A.521213",
                "2906.BMA.006.005.A.521213",
                "2907.BMA.006.005.A.500001"
            ]
        }
    }

    # 3️⃣ Apply table replacement
    insert_tables_from_markers(doc, tables)

    # 4️⃣ Save output
    doc.save("output_test.docx")

    print("✅ Test completed. Open output_test.docx")

def generate_documents_from_json(
    json_path: str,
    output_dir: str = "."
):
    """
    Generate DOCX files from a JSON job definition.

    Args:
        json_path (str): Path to data.json
        output_dir (str): Directory to save generated DOCX files
    """

    with open(json_path, "r", encoding="utf-8") as f:
        jobs = json.load(f)

    for job in jobs:
        doc = Document("template/"+job["template"])

        include_blocks(doc, job.get("includedBlocks", []))
        replace_placeholders(doc, job.get("replacements", {}))

        output_path = f'{output_dir}/{job["filename"]}.docx'
        doc.save(output_path)

        print("Generated:", output_path)

def main():
    template = sys.argv[1]
    output = sys.argv[2]
    replacements = json.loads(sys.argv[3])
    included_blocks_arg = sys.argv[4]
    datatables_arg = sys.argv[5]

    doc = Document(template)

    # # ✅ skip include_blocks if argv[4] == "null"
    if included_blocks_arg != "null":
        included_blocks = json.loads(included_blocks_arg)
        include_blocks(doc, included_blocks)
    else:
        included_blocks = []

    combined_replacements = build_combined_replacements(replacements, included_blocks)
    replace_placeholders(doc, combined_replacements)
    replace_placeholders_in_headers_and_footers(doc, combined_replacements)


    # ✅ skip datatables if argv[5] == "null"
    if datatables_arg != "null":
        datatables = json.loads(datatables_arg)
        insert_tables_from_markers(doc, datatables)

    fix_footer_first_page_only(doc)

    doc.save(output)

if __name__ == "__main__":
    main()
    # generate_documents_from_json("data.example.json")
    # test_insert_tables()


